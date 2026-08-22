from flask import (
    Blueprint,
    render_template,
    redirect,
    url_for,
    flash,
    request,
    jsonify
)
from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    flash
)

from models import db

import io

from flask import send_file

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Alignment

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape
from reportlab.lib.pagesizes import A4

from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle
)

import csv
import io

from flask import jsonify

from flask import send_file

from models.faculty import Faculty
from models.subject import Subject
from models.semester import Semester
from models.section import Section
from models.classroom import Classroom
from models.department import Department
from models.timeslot import TimeSlot
from models.timetable import Timetable
from models.subject_allocation import SubjectAllocation
from models.faculty_availability import FacultyAvailability


from ga.genetic_algorithm import GeneticAlgorithm


timetable_bp = Blueprint(
    "timetable",
    __name__,
    url_prefix="/timetable"
)


# =====================================================
# VALID LAB PERIOD PAIRS
# =====================================================

VALID_LAB_START_PERIODS = [
    1,
    3,
    5,
    6
]

# =========================================================
# MAIN TIMETABLE PAGE
# =========================================================


@timetable_bp.route("/")
def timetable():

    # =====================================================
    # SELECTED FILTER VALUES
    # =====================================================

    selected_department = request.args.get(
        "department",
        ""
    ).strip()

    selected_semester = request.args.get(
        "semester",
        ""
    ).strip()

    selected_section = request.args.get(
        "section",
        ""
    ).strip()

    # =====================================================
    # LOAD UNIQUE DEPARTMENTS
    # =====================================================

    departments = (
        db.session.query(
            Department.department_name
        )
        .distinct()
        .order_by(
            Department.department_name
        )
        .all()
    )

    # =====================================================
    # LOAD UNIQUE SEMESTERS
    # =====================================================

    semesters = (
        db.session.query(
            Semester.semester
        )
        .distinct()
        .order_by(
            Semester.semester
        )
        .all()
    )

    # =====================================================
    # LOAD UNIQUE SECTIONS
    # =====================================================

    sections = (
        db.session.query(
            Section.section_name
        )
        .distinct()
        .order_by(
            Section.section_name
        )
        .all()
    )

    # =====================================================
    # LOAD TIMETABLE
    # =====================================================

    lectures = []

    if (
        selected_department
        and selected_semester
        and selected_section
    ):

        try:

            semester_number = int(
                selected_semester
            )

        except ValueError:

            semester_number = None

        if semester_number is not None:

            lectures = (
                Timetable.query
                .filter(
                    Timetable.department
                    == selected_department,

                    Timetable.semester
                    == semester_number,

                    Timetable.section
                    == selected_section
                )
                .order_by(
                    Timetable.day,
                    Timetable.period
                )
                .all()
            )

        # =====================================================
        # HEADER TIME SLOTS
        # =====================================================

    header_slots = (
        TimeSlot.query
        .filter_by(day="Monday")
        .order_by(
            db.case(
                (TimeSlot.period == 1, 1),
                (TimeSlot.period == 2, 2),
                (TimeSlot.session == "Break", 3),
                (TimeSlot.period == 3, 4),
                (TimeSlot.period == 4, 5),
                (TimeSlot.session == "Lunch", 6),
                (TimeSlot.period == 5, 7),
                (TimeSlot.period == 6, 8),
                (TimeSlot.period == 7, 9)
            )
        )
        .all()
    )

    # =====================================================
    # RENDER PAGE
    # =====================================================

    return render_template(
        "timetable/timetable.html",

        departments=departments,

        semesters=semesters,

        sections=sections,

        selected_department=(
            selected_department
        ),

        selected_semester=(
            selected_semester
        ),

        selected_section=(
            selected_section
        ),

        lectures=lectures,

        timetables=lectures,

        header_slots=header_slots,

        department=selected_department,

        semester=selected_semester,

        section=selected_section
    )

# =========================================================
# GET SEMESTERS
# =========================================================


@timetable_bp.route("/get_semesters")
def get_semesters():

    from flask import request

    department = request.args.get("department")

    semesters = (
        db.session.query(Timetable.semester)
        .filter(Timetable.department == department)
        .distinct()
        .order_by(Timetable.semester)
        .all()
    )

    return jsonify(
        [row[0] for row in semesters]
    )

# =========================================================
# GET SECTIONS
# =========================================================


@timetable_bp.route("/get_sections")
def get_sections():

    from flask import request

    department = request.args.get("department")

    semester = request.args.get("semester")

    sections = (

        db.session.query(Timetable.section)

        .filter(

            Timetable.department == department,

            Timetable.semester == semester

        )

        .distinct()

        .order_by(Timetable.section)

        .all()

    )

    return jsonify(

        [row[0] for row in sections]

    )


# =========================================================
# FACULTY TIMETABLE
# =========================================================

@timetable_bp.route("/faculty")
def faculty_timetable():

    # -----------------------------------------------------
    # GET ALL FACULTY
    # -----------------------------------------------------

    faculty_records = (
        Faculty.query
        .order_by(Faculty.faculty_name)
        .all()
    )

    faculty_names = [
        faculty.faculty_name
        for faculty in faculty_records
        if faculty.faculty_name
    ]

    # -----------------------------------------------------
    # SELECTED FACULTY
    # -----------------------------------------------------

    selected_faculty = request.args.get(
        "faculty",
        ""
    ).strip()

    timetables = []

    # -----------------------------------------------------
    # GET SELECTED FACULTY TIMETABLE
    # -----------------------------------------------------

    if selected_faculty:

        timetables = (
            Timetable.query
            .filter(
                Timetable.faculty == selected_faculty
            )
            .all()
        )

        day_order = {
            "Monday": 1,
            "Tuesday": 2,
            "Wednesday": 3,
            "Thursday": 4,
            "Friday": 5
        }

        timetables.sort(
            key=lambda row: (
                day_order.get(
                    row.day,
                    99
                ),
                row.period
            )
        )

    # -----------------------------------------------------
    # HEADER SLOTS
    # -----------------------------------------------------

    header_slots = (
        TimeSlot.query
        .filter_by(day="Monday")
        .order_by(
            TimeSlot.period
        )
        .all()
    )

    # -----------------------------------------------------
    # RENDER
    # -----------------------------------------------------

    return render_template(
        "timetable/faculty_timetable.html",
        faculty_names=faculty_names,
        selected_faculty=selected_faculty,
        timetables=timetables,
        header_slots=header_slots
    )


# =========================================================
# GENERATE TIMETABLE
# =========================================================

@timetable_bp.route("/generate", methods=["POST"])
def generate_timetable():

    try:

        # =====================================================
        # LOAD DATA
        # =====================================================

        faculty = Faculty.query.all()

        subjects = Subject.query.all()

        classrooms = Classroom.query.all()

        departments = Department.query.all()

        timeslots = (
            TimeSlot.query
            .filter(
                TimeSlot.period > 0
            )
            .order_by(
                TimeSlot.day,
                TimeSlot.period
            )
            .all()
        )

        allocations = (
            SubjectAllocation.query
            .all()
        )

        availability = (
            FacultyAvailability.query
            .all()
        )

        # =====================================================
        # CHECK REQUIRED DATA
        # =====================================================

        if not allocations:

            flash(
                "No subject allocations found.",
                "warning"
            )

            return redirect(
                url_for(
                    "timetable.generate_page"
                )
            )

        if not timeslots:

            flash(
                "No teaching time slots found.",
                "warning"
            )

            return redirect(
                url_for(
                    "timetable.generate_page"
                )
            )

        if not classrooms:

            flash(
                "No classrooms found.",
                "warning"
            )

            return redirect(
                url_for(
                    "timetable.generate_page"
                )
            )

        # =====================================================
        # DEBUG
        # =====================================================

        print()
        print("==========================================")
        print("STARTING TIMETABLE GENERATION")
        print("==========================================")
        print("Faculty:", len(faculty))
        print("Subjects:", len(subjects))
        print("Classrooms:", len(classrooms))
        print("Departments:", len(departments))
        print("TimeSlots:", len(timeslots))
        print("Allocations:", len(allocations))
        print("==========================================")
        print()

        # =====================================================
        # CREATE GA
        # =====================================================

        ga = GeneticAlgorithm(

            faculty=faculty,

            subjects=subjects,

            classrooms=classrooms,

            departments=departments,

            timeslots=timeslots,

            allocations=allocations,

            availability=availability,

            population_size=30,

            generations=100,

            mutation_rate=0.10
        )

        # =====================================================
        # RUN GA
        # =====================================================

        best_solution = ga.run()

        if not best_solution:

            raise Exception(
                "Genetic Algorithm could not generate "
                "a valid timetable."
            )

        # =====================================================
        # GET GENERATED LECTURES
        # =====================================================

        if hasattr(
            best_solution,
            "timetable"
        ):

            generated_timetable = (
                best_solution.timetable
            )

        else:

            generated_timetable = (
                best_solution
            )

        if not generated_timetable:

            raise Exception(
                "Genetic Algorithm returned "
                "an empty timetable."
            )

        # =====================================================
        # NORMALIZE GENERATED DATA FIRST
        #
        # IMPORTANT:
        # Nothing is deleted from the database before this
        # section has successfully completed.
        # =====================================================

        prepared_entries = []

        affected_scopes = set()

        for lecture in generated_timetable:

            # =================================================
            # GET IDS
            # =================================================

            allocation_id = lecture.get(
                "allocation_id"
            )

            subject_id = lecture.get(
                "subject_id"
            )

            faculty_id = lecture.get(
                "faculty_id"
            )

            room_id = lecture.get(
                "room_id",
                lecture.get(
                    "classroom_id"
                )
            )

            timeslot_id = lecture.get(
                "timeslot_id"
            )

            # =================================================
            # DEPARTMENT
            # =================================================

            department = str(
                lecture.get(
                    "department",
                    ""
                )
            ).strip()

            if not department:

                department_id = lecture.get(
                    "department_id"
                )

                if department_id:

                    department_obj = (
                        db.session.get(
                            Department,
                            int(department_id)
                        )
                    )

                    if department_obj:

                        department = str(
                            department_obj.department_name
                        ).strip()

            # =================================================
            # FALLBACK: GET DEPARTMENT FROM ALLOCATION
            # =================================================

            if not department and allocation_id:

                allocation_obj = (
                    db.session.get(
                        SubjectAllocation,
                        int(allocation_id)
                    )
                )

                if allocation_obj:

                    # Try common possible department fields
                    allocation_department = getattr(
                        allocation_obj,
                        "department",
                        None
                    )

                    if allocation_department:

                        if isinstance(
                            allocation_department,
                            str
                        ):

                            department = (
                                allocation_department.strip()
                            )

                        else:

                            department = str(
                                getattr(
                                    allocation_department,
                                    "department_name",
                                    ""
                                )
                            ).strip()

                    if not department:

                        allocation_department_id = (
                            getattr(
                                allocation_obj,
                                "department_id",
                                None
                            )
                        )

                        if allocation_department_id:

                            department_obj = (
                                db.session.get(
                                    Department,
                                    int(
                                        allocation_department_id
                                    )
                                )
                            )

                            if department_obj:

                                department = str(
                                    department_obj.department_name
                                ).strip()

            # =================================================
            # SEMESTER
            # =================================================

            semester_value = lecture.get(
                "semester"
            )

            semester = 0

            if isinstance(
                semester_value,
                int
            ):

                semester = semester_value

            elif isinstance(
                semester_value,
                str
            ):

                try:

                    semester = int(
                        semester_value
                    )

                except (
                    ValueError,
                    TypeError
                ):

                    semester = 0

            elif hasattr(
                semester_value,
                "semester"
            ):

                semester = int(
                    semester_value.semester
                )

            else:

                semester_id = lecture.get(
                    "semester_id"
                )

                if semester_id:

                    semester_obj = (
                        db.session.get(
                            Semester,
                            int(semester_id)
                        )
                    )

                    if semester_obj:

                        semester = int(
                            semester_obj.semester
                        )

            # =================================================
            # FALLBACK: SEMESTER FROM ALLOCATION
            # =================================================

            if not semester and allocation_id:

                allocation_obj = (
                    db.session.get(
                        SubjectAllocation,
                        int(allocation_id)
                    )
                )

                if allocation_obj:

                    allocation_semester = getattr(
                        allocation_obj,
                        "semester",
                        None
                    )

                    if isinstance(
                        allocation_semester,
                        int
                    ):

                        semester = (
                            allocation_semester
                        )

                    elif allocation_semester:

                        semester = int(
                            getattr(
                                allocation_semester,
                                "semester",
                                0
                            )
                        )

                    if not semester:

                        semester_id = getattr(
                            allocation_obj,
                            "semester_id",
                            None
                        )

                        if semester_id:

                            semester_obj = (
                                db.session.get(
                                    Semester,
                                    int(
                                        semester_id
                                    )
                                )
                            )

                            if semester_obj:

                                semester = int(
                                    semester_obj.semester
                                )

            # =================================================
            # SECTION
            # =================================================

            section = str(
                lecture.get(
                    "section",
                    ""
                )
            ).strip()

            if not section:

                section_id = lecture.get(
                    "section_id"
                )

                if section_id:

                    section_obj = (
                        db.session.get(
                            Section,
                            int(section_id)
                        )
                    )

                    if section_obj:

                        section = str(
                            section_obj.section_name
                        ).strip()

            # =================================================
            # FALLBACK: SECTION FROM ALLOCATION
            # =================================================

            if not section and allocation_id:

                allocation_obj = (
                    db.session.get(
                        SubjectAllocation,
                        int(allocation_id)
                    )
                )

                if allocation_obj:

                    allocation_section = getattr(
                        allocation_obj,
                        "section",
                        None
                    )

                    if allocation_section:

                        if isinstance(
                            allocation_section,
                            str
                        ):

                            section = (
                                allocation_section.strip()
                            )

                        else:

                            section = str(
                                getattr(
                                    allocation_section,
                                    "section_name",
                                    ""
                                )
                            ).strip()

                    if not section:

                        allocation_section_id = (
                            getattr(
                                allocation_obj,
                                "section_id",
                                None
                            )
                        )

                        if allocation_section_id:

                            section_obj = (
                                db.session.get(
                                    Section,
                                    int(
                                        allocation_section_id
                                    )
                                )
                            )

                            if section_obj:

                                section = str(
                                    section_obj.section_name
                                ).strip()

            # =================================================
            # SUBJECT
            # =================================================

            subject = str(
                lecture.get(
                    "subject",
                    ""
                )
            ).strip()

            if not subject and subject_id:

                subject_obj = (
                    db.session.get(
                        Subject,
                        int(subject_id)
                    )
                )

                if subject_obj:

                    subject = str(
                        subject_obj.subject_name
                    ).strip()

            # =================================================
            # FACULTY
            # =================================================

            faculty_name = str(
                lecture.get(
                    "faculty",
                    ""
                )
            ).strip()

            if not faculty_name and faculty_id:

                faculty_obj = (
                    db.session.get(
                        Faculty,
                        int(faculty_id)
                    )
                )

                if faculty_obj:

                    faculty_name = str(
                        faculty_obj.faculty_name
                    ).strip()

            # =================================================
            # ROOM
            # =================================================

            room = str(
                lecture.get(
                    "room",
                    ""
                )
            ).strip()

            if not room and room_id:

                room_obj = (
                    db.session.get(
                        Classroom,
                        int(room_id)
                    )
                )

                if room_obj:

                    room = str(
                        room_obj.room_name
                    ).strip()

            # =================================================
            # SUBJECT TYPE
            # =================================================

            subject_type = str(

                lecture.get(
                    "subject_type",
                    lecture.get(
                        "type",
                        "Theory"
                    )
                )

            ).strip()

            # =================================================
            # DAY
            # =================================================

            day = str(
                lecture.get(
                    "day",
                    ""
                )
            ).strip()

            # =================================================
            # PERIOD
            # =================================================

            try:

                period = int(
                    lecture.get(
                        "period",
                        0
                    ) or 0
                )

            except (
                ValueError,
                TypeError
            ):

                period = 0

            # =================================================
            # TIMES
            # =================================================

            start_time = str(
                lecture.get(
                    "start_time",
                    ""
                )
            ).strip()

            end_time = str(
                lecture.get(
                    "end_time",
                    ""
                )
            ).strip()

            session = str(
                lecture.get(
                    "session",
                    "Teaching"
                )
            ).strip()

            # =================================================
            # GET TIMES FROM TIMESLOT IF MISSING
            # =================================================

            if (
                timeslot_id
                and (
                    not start_time
                    or not end_time
                )
            ):

                timeslot_obj = (
                    db.session.get(
                        TimeSlot,
                        int(timeslot_id)
                    )
                )

                if timeslot_obj:

                    start_time = str(
                        timeslot_obj.start_time
                    )

                    end_time = str(
                        timeslot_obj.end_time
                    )

                    if not day:

                        day = str(
                            timeslot_obj.day
                        ).strip()

            # =================================================
            # VALIDATE
            # =================================================

            if not all([

                department,

                semester,

                section,

                subject,

                faculty_name,

                room,

                day,

                period,

                start_time,

                end_time

            ]):

                print(
                    "SKIPPED INCOMPLETE:",
                    lecture
                )

                continue

            # =================================================
            # RECORD AFFECTED SCOPE
            # =================================================

            affected_scopes.add(
                (
                    department,
                    int(semester),
                    section
                )
            )

            # =================================================
            # STORE PREPARED ENTRY
            #
            # DO NOT INSERT YET.
            # =================================================

            prepared_entries.append({

                "department": department,

                "semester": int(
                    semester
                ),

                "section": section,

                "subject": subject,

                "subject_type": subject_type,

                "faculty": faculty_name,

                "room": room,

                "day": day,

                "period": period,

                "start_time": start_time,

                "end_time": end_time,

                "session": session,

                "allocation_id": int(
                    allocation_id or 0
                ),

                "timeslot_id": int(
                    timeslot_id or 0
                )

            })

        # =====================================================
        # MAKE SURE SOMETHING VALID WAS GENERATED
        # =====================================================

        if not prepared_entries:

            db.session.rollback()

            flash(
                "GA generated data, but no valid "
                "classes could be saved.",
                "danger"
            )

            return redirect(
                url_for(
                    "timetable.generate_page"
                )
            )

        # =====================================================
        # DEBUG AFFECTED SCOPES
        # =====================================================

        print()
        print("==========================================")
        print("PREPARED TIMETABLE")
        print("==========================================")
        print(
            "Valid Classes:",
            len(prepared_entries)
        )
        print(
            "Affected Scopes:",
            sorted(
                affected_scopes
            )
        )
        print("==========================================")
        print()

        # =====================================================
        # CLEAR OLD TIMETABLE
        #
        # THIS IS NOW SAFE:
        # department / semester / section are already known.
        # =====================================================

        deleted_count = 0

        for (
            department,
            semester,
            section
        ) in affected_scopes:

            deleted = (
                Timetable.query
                .filter_by(
                    department=department,
                    semester=semester,
                    section=section
                )
                .delete(
                    synchronize_session=False
                )
            )

            deleted_count += deleted

        # =====================================================
        # SAVE NEW TIMETABLE
        # =====================================================

        created_count = 0

        first_department = None
        first_semester = None
        first_section = None

        for entry in prepared_entries:

            timetable_entry = Timetable(

                department=entry[
                    "department"
                ],

                semester=entry[
                    "semester"
                ],

                section=entry[
                    "section"
                ],

                subject=entry[
                    "subject"
                ],

                subject_type=entry[
                    "subject_type"
                ],

                faculty=entry[
                    "faculty"
                ],

                room=entry[
                    "room"
                ],

                day=entry[
                    "day"
                ],

                period=entry[
                    "period"
                ],

                start_time=entry[
                    "start_time"
                ],

                end_time=entry[
                    "end_time"
                ],

                session=entry[
                    "session"
                ],

                allocation_id=entry[
                    "allocation_id"
                ],

                timeslot_id=entry[
                    "timeslot_id"
                ]

            )

            db.session.add(
                timetable_entry
            )

            if first_department is None:

                first_department = entry[
                    "department"
                ]

                first_semester = entry[
                    "semester"
                ]

                first_section = entry[
                    "section"
                ]

            created_count += 1

        # =====================================================
        # SINGLE COMMIT
        # =====================================================

        db.session.commit()

        # =====================================================
        # SUCCESS
        # =====================================================

        print()
        print("==========================================")
        print("TIMETABLE GENERATED SUCCESSFULLY")
        print("==========================================")
        print(
            "Old Classes Deleted:",
            deleted_count
        )
        print(
            "New Classes Created:",
            created_count
        )
        print(
            "Department:",
            first_department
        )
        print(
            "Semester:",
            first_semester
        )
        print(
            "Section:",
            first_section
        )
        print("==========================================")
        print()

        flash(

            f"Timetable generated successfully! "
            f"{created_count} classes created.",

            "success"

        )

        # =====================================================
        # REDIRECT
        # =====================================================

        return redirect(

            url_for(

                "timetable.timetable",

                department=first_department,

                semester=first_semester,

                section=first_section

            )

        )

    except Exception as e:

        db.session.rollback()

        print()
        print("==========================================")
        print("TIMETABLE GENERATION ERROR")
        print("==========================================")
        print(
            "ERROR TYPE:",
            type(e).__name__
        )
        print(
            "ERROR:",
            str(e)
        )
        print("==========================================")
        print()

        flash(

            f"Timetable generation failed: {str(e)}",

            "danger"

        )

        return redirect(

            url_for(
                "timetable.generate_page"
            )

        )


# =========================================================
# CLEAR TIMETABLE
# =========================================================


@timetable_bp.route("/clear", methods=["POST"])
def clear_timetable():

    Timetable.query.delete()

    db.session.commit()

    flash(

        "Timetable cleared successfully.",

        "success"

    )

    return redirect(
        url_for("timetable.timetable")
    )


# =========================================================
# REPORTS
# =========================================================
@timetable_bp.route("/reports")
def reports():

    return render_template(
        "timetable/reports.html"
    )

# =========================================================
# EXPORT EXCEL
# =========================================================


@timetable_bp.route("/export/excel")
def export_excel():

    timetables = (
        Timetable.query
        .order_by(
            Timetable.department,
            Timetable.semester,
            Timetable.section,
            Timetable.day,
            Timetable.period
        )
        .all()
    )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "College Timetable"

    headers = [

        "Department",

        "Semester",

        "Section",

        "Day",

        "Period",

        "Start Time",

        "End Time",

        "Subject",

        "Faculty",

        "Room",

        "Type"

    ]

    header_fill = PatternFill(
        fill_type="solid",
        start_color="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    center = Alignment(
        horizontal="center",
        vertical="center"
    )

    # Header Row

    for col, header in enumerate(headers, start=1):

        cell = sheet.cell(
            row=1,
            column=col
        )

        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    # Data Rows

    row_no = 2

    for item in timetables:

        sheet.cell(row=row_no, column=1).value = item.department
        sheet.cell(row=row_no, column=2).value = item.semester
        sheet.cell(row=row_no, column=3).value = item.section
        sheet.cell(row=row_no, column=4).value = item.day
        sheet.cell(row=row_no, column=5).value = item.period
        sheet.cell(row=row_no, column=6).value = item.start_time
        sheet.cell(row=row_no, column=7).value = item.end_time
        sheet.cell(row=row_no, column=8).value = item.subject
        sheet.cell(row=row_no, column=9).value = item.faculty
        sheet.cell(row=row_no, column=10).value = item.room
        sheet.cell(row=row_no, column=11).value = item.subject_type

        row_no += 1

    # Auto Width

    for column in sheet.columns:

        length = 0

        column_letter = column[0].column_letter

        for cell in column:

            try:

                if len(str(cell.value)) > length:

                    length = len(str(cell.value))

            except:

                pass

        sheet.column_dimensions[column_letter].width = length + 5

    output = io.BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(

        output,

        download_name="College_Timetable.xlsx",

        as_attachment=True,

        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    )


# =========================================================
# EXPORT PDF
# =========================================================

@timetable_bp.route("/export/pdf")
def export_pdf():

    timetables = (
        Timetable.query
        .order_by(
            Timetable.department,
            Timetable.semester,
            Timetable.section,
            Timetable.day,
            Timetable.period
        )
        .all()
    )

    buffer = io.BytesIO()

    pdf = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4)
    )

    data = [[
        "Department",
        "Semester",
        "Section",
        "Day",
        "Period",
        "Subject",
        "Faculty",
        "Room"
    ]]

    for row in timetables:

        data.append([

            row.department,

            row.semester,

            row.section,

            row.day,

            row.period,

            row.subject,

            row.faculty,

            row.room

        ])

    table = Table(data)

    table.setStyle(

        TableStyle([

            ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),

            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),

            ("GRID", (0, 0), (-1, -1), 1, colors.black),

            ("BACKGROUND", (0, 1), (-1, -1), colors.beige),

            ("ALIGN", (0, 0), (-1, -1), "CENTER"),

            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),

            ("BOTTOMPADDING", (0, 0), (-1, 0), 10)

        ])

    )

    pdf.build([table])

    buffer.seek(0)

    return send_file(

        buffer,

        as_attachment=True,

        download_name="College_Timetable.pdf",

        mimetype="application/pdf"

    )


# =========================================================
# EXPORT CSV
# =========================================================

@timetable_bp.route("/export/csv")
def export_csv():

    timetables = (
        Timetable.query
        .order_by(
            Timetable.department,
            Timetable.semester,
            Timetable.section,
            Timetable.day,
            Timetable.period
        )
        .all()
    )

    output = io.StringIO()

    writer = csv.writer(output)

    writer.writerow([
        "Department",
        "Semester",
        "Section",
        "Day",
        "Period",
        "Start Time",
        "End Time",
        "Subject",
        "Faculty",
        "Room",
        "Subject Type"
    ])

    for row in timetables:

        writer.writerow([

            row.department,

            row.semester,

            row.section,

            row.day,

            row.period,

            row.start_time,

            row.end_time,

            row.subject,

            row.faculty,

            row.room,

            row.subject_type

        ])

    memory_file = io.BytesIO()

    memory_file.write(
        output.getvalue().encode("utf-8")
    )

    memory_file.seek(0)

    output.close()

    return send_file(

        memory_file,

        mimetype="text/csv",

        as_attachment=True,

        download_name="College_Timetable.csv"

    )
# =========================================================
# GENERATE PAGE
# =========================================================


@timetable_bp.route("/generate-page")
def generate_page():

    departments = (
        Department.query
        .order_by(
            Department.department_name
        )
        .all()
    )

    semesters = (
        Semester.query
        .order_by(
            Semester.semester
        )
        .all()
    )

    sections = (
        Section.query
        .order_by(
            Section.section_name
        )
        .all()
    )

    return render_template(
        "timetable/generate.html",
        departments=departments,
        semesters=semesters,
        sections=sections
    )
